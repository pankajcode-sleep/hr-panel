from flask import Blueprint, render_template, redirect, url_for, flash, request, send_file, jsonify
from flask_login import login_required, current_user
from datetime import date, datetime, time as time_type
from sqlalchemy import extract
import calendar
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import io

from app.extensions import db
from app.models.attendance import Attendance
from app.models.employee import Employee
from app.utils.decorators import hr_required

attendance_bp = Blueprint('attendance', __name__)

PUNCH_PHOTO_FOLDER = os.path.join('app', 'static', 'uploads', 'punch_photos')
os.makedirs(PUNCH_PHOTO_FOLDER, exist_ok=True)


@attendance_bp.route('/punch')
@login_required
def punch():
    employee = Employee.query.filter_by(user_id=current_user.id).first()
    today = date.today()
    today_record = None
    if employee:
        today_record = Attendance.query.filter_by(
            employee_id=employee.id, date=today
        ).first()
    return render_template('attendance/punch.html',
                           employee=employee,
                           today_record=today_record,
                           today=today)


@attendance_bp.route('/punch-in', methods=['POST'])
@login_required
def punch_in():
    import base64
    employee = Employee.query.filter_by(user_id=current_user.id).first()
    if not employee:
        return jsonify({'success': False, 'message': 'Employee not linked!'})

    today = date.today()
    now = datetime.now()
    existing = Attendance.query.filter_by(
        employee_id=employee.id, date=today
    ).first()

    if existing and existing.check_in:
        return jsonify({'success': False, 'message': 'Already punched in!'})

    lat = request.form.get('latitude')
    lng = request.form.get('longitude')
    address = request.form.get('address', '')
    photo_data = request.form.get('photo')

    photo_filename = None
    if photo_data:
        try:
            photo_bytes = base64.b64decode(photo_data.split(',')[1])
            photo_filename = f'punch_in_{employee.employee_id}_{today}.jpg'
            with open(os.path.join(PUNCH_PHOTO_FOLDER, photo_filename), 'wb') as f:
                f.write(photo_bytes)
        except:
            photo_filename = None

    check_in_time = now.time().replace(second=0, microsecond=0)
    late_minutes = Attendance.calculate_late_minutes(check_in_time)
    status = 'Late' if late_minutes > Attendance.LATE_THRESHOLD else 'Present'

    if existing:
        existing.check_in = check_in_time
        existing.punch_in_lat = lat
        existing.punch_in_lng = lng
        existing.punch_in_address = address
        existing.punch_in_photo = photo_filename
        existing.late_minutes = late_minutes
        existing.status = status
    else:
        att = Attendance(
            employee_id=employee.id,
            date=today,
            check_in=check_in_time,
            status=status,
            late_minutes=late_minutes,
            punch_in_lat=lat,
            punch_in_lng=lng,
            punch_in_address=address,
            punch_in_photo=photo_filename,
            marked_by=current_user.id
        )
        db.session.add(att)

    db.session.commit()
    return jsonify({
        'success': True,
        'message': f'Punched In at {check_in_time.strftime("%I:%M %p")}',
        'status': status,
        'time': check_in_time.strftime('%I:%M %p'),
        'late_minutes': late_minutes
    })


@attendance_bp.route('/punch-out', methods=['POST'])
@login_required
def punch_out():
    import base64
    employee = Employee.query.filter_by(user_id=current_user.id).first()
    if not employee:
        return jsonify({'success': False, 'message': 'Employee not linked!'})

    today = date.today()
    now = datetime.now()
    attendance = Attendance.query.filter_by(
        employee_id=employee.id, date=today
    ).first()

    if not attendance or not attendance.check_in:
        return jsonify({'success': False, 'message': 'Please punch in first!'})
    if attendance.check_out:
        return jsonify({'success': False, 'message': 'Already punched out!'})

    lat = request.form.get('latitude')
    lng = request.form.get('longitude')
    address = request.form.get('address', '')
    photo_data = request.form.get('photo')

    photo_filename = None
    if photo_data:
        try:
            photo_bytes = base64.b64decode(photo_data.split(',')[1])
            photo_filename = f'punch_out_{employee.employee_id}_{today}.jpg'
            with open(os.path.join(PUNCH_PHOTO_FOLDER, photo_filename), 'wb') as f:
                f.write(photo_bytes)
        except:
            photo_filename = None

    check_out_time = now.time().replace(second=0, microsecond=0)
    attendance.check_out = check_out_time
    attendance.punch_out_lat = lat
    attendance.punch_out_lng = lng
    attendance.punch_out_address = address
    attendance.punch_out_photo = photo_filename
    db.session.commit()

    return jsonify({
        'success': True,
        'message': f'Punched Out at {check_out_time.strftime("%I:%M %p")}',
        'time': check_out_time.strftime('%I:%M %p'),
        'working_hours': attendance.working_hours
    })


@attendance_bp.route('/punch-records')
@hr_required
def punch_records():
    selected_date = request.args.get('date', str(date.today()))
    try:
        selected_date = date.fromisoformat(selected_date)
    except:
        selected_date = date.today()
    records = Attendance.query.filter_by(
        date=selected_date
    ).order_by(Attendance.check_in).all()
    return render_template('attendance/punch_records.html',
                           records=records,
                           selected_date=selected_date)


@attendance_bp.route('/mark', methods=['GET', 'POST'])
@hr_required
def mark():
    today = date.today()
    selected_date = request.args.get('date', str(today))
    try:
        selected_date = date.fromisoformat(selected_date)
    except:
        selected_date = today

    employees = Employee.query.filter_by(status='Active').order_by(
        Employee.first_name
    ).all()
    existing = {}
    records = Attendance.query.filter_by(date=selected_date).all()
    for rec in records:
        existing[rec.employee_id] = rec

    if request.method == 'POST':
        for emp in employees:
            status = request.form.get(f'status_{emp.id}')
            check_in_str = request.form.get(f'check_in_{emp.id}')
            check_out_str = request.form.get(f'check_out_{emp.id}')
            note = request.form.get(f'note_{emp.id}')
            if not status:
                continue
            check_in = None
            check_out = None
            late_minutes = 0
            if check_in_str:
                try:
                    check_in = time_type.fromisoformat(check_in_str)
                    late_minutes = Attendance.calculate_late_minutes(check_in)
                    if late_minutes > Attendance.LATE_THRESHOLD and status == 'Present':
                        status = 'Late'
                except:
                    pass
            if check_out_str:
                try:
                    check_out = time_type.fromisoformat(check_out_str)
                except:
                    pass
            if emp.id in existing:
                att = existing[emp.id]
                att.status = status
                att.check_in = check_in
                att.check_out = check_out
                att.late_minutes = late_minutes
                att.note = note
                att.marked_by = current_user.id
            else:
                att = Attendance(
                    employee_id=emp.id,
                    date=selected_date,
                    status=status,
                    check_in=check_in,
                    check_out=check_out,
                    late_minutes=late_minutes,
                    note=note,
                    marked_by=current_user.id
                )
                db.session.add(att)
        db.session.commit()
        flash(f'Attendance saved for {selected_date}!', 'success')
        return redirect(url_for('attendance.mark', date=selected_date))

    return render_template('attendance/mark.html',
                           employees=employees,
                           existing=existing,
                           selected_date=selected_date,
                           today=today)


@attendance_bp.route('/monthly-report')
@hr_required
def monthly_report():
    month = request.args.get('month', date.today().month, type=int)
    year = request.args.get('year', date.today().year, type=int)
    _, days_in_month = calendar.monthrange(year, month)
    month_dates = [date(year, month, d) for d in range(1, days_in_month + 1)]
    working_days = [d for d in month_dates if d.weekday() != 6]
    employees = Employee.query.filter_by(status='Active').order_by(
        Employee.first_name
    ).all()
    report_data = []
    for emp in employees:
        records = Attendance.query.filter(
            Attendance.employee_id == emp.id,
            extract('month', Attendance.date) == month,
            extract('year', Attendance.date) == year
        ).all()
        summary = {'Present': 0, 'Absent': 0, 'Late': 0,
                   'Half Day': 0, 'WFH': 0, 'Holiday': 0}
        total_late_minutes = 0
        for rec in records:
            if rec.status in summary:
                summary[rec.status] += 1
            total_late_minutes += rec.late_minutes or 0
        marked_dates = {rec.date for rec in records}
        unmarked = len([d for d in working_days if d not in marked_dates])
        summary['Absent'] += unmarked
        attendance_percent = round(
            (summary['Present'] + summary['Late'] + summary['WFH'] +
             summary['Half Day'] * 0.5) / len(working_days) * 100, 1
        ) if working_days else 0
        report_data.append({
            'employee': emp,
            'summary': summary,
            'total_late_minutes': total_late_minutes,
            'working_days': len(working_days),
            'attendance_percent': attendance_percent
        })
    return render_template('attendance/monthly_report.html',
                           report_data=report_data,
                           month=month, year=year,
                           month_name=calendar.month_name[month],
                           working_days=len(working_days))


@attendance_bp.route('/employee/<int:id>')
@login_required
def employee_attendance(id):
    employee = Employee.query.get_or_404(id)
    month = request.args.get('month', date.today().month, type=int)
    year = request.args.get('year', date.today().year, type=int)
    _, days_in_month = calendar.monthrange(year, month)
    records = Attendance.query.filter(
        Attendance.employee_id == id,
        extract('month', Attendance.date) == month,
        extract('year', Attendance.date) == year
    ).order_by(Attendance.date).all()
    records_dict = {rec.date: rec for rec in records}
    all_dates = []
    for day in range(1, days_in_month + 1):
        d = date(year, month, day)
        rec = records_dict.get(d)
        all_dates.append({
            'date': d, 'record': rec,
            'is_sunday': d.weekday() == 6,
            'is_today': d == date.today()
        })
    summary = {
        'Present': len([r for r in records if r.status == 'Present']),
        'Absent': len([r for r in records if r.status == 'Absent']),
        'Late': len([r for r in records if r.status == 'Late']),
        'Half Day': len([r for r in records if r.status == 'Half Day']),
        'WFH': len([r for r in records if r.status == 'WFH']),
    }
    return render_template('attendance/employee_attendance.html',
                           employee=employee,
                           all_dates=all_dates,
                           summary=summary,
                           month=month, year=year,
                           month_name=calendar.month_name[month])


@attendance_bp.route('/export/<int:month>/<int:year>')
@hr_required
def export_attendance(month, year):
    employees = Employee.query.filter_by(status='Active').order_by(
        Employee.first_name
    ).all()
    _, days_in_month = calendar.monthrange(year, month)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f'Attendance {calendar.month_name[month]} {year}'
    dark = PatternFill(start_color='1a1a2e', end_color='1a1a2e', fill_type='solid')
    white_font = Font(color='FFFFFF', bold=True)
    center = Alignment(horizontal='center')
    headers = ['Employee', 'ID', 'Dept'] + \
              [str(d) for d in range(1, days_in_month + 1)] + \
              ['P', 'A', 'L', 'H', 'W', '%']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = dark
        cell.font = white_font
        cell.alignment = center
    status_short = {
        'Present': 'P', 'Absent': 'A', 'Late': 'L',
        'Half Day': 'H', 'WFH': 'W', 'Holiday': 'Ho'
    }
    for row_num, emp in enumerate(employees, 2):
        records = Attendance.query.filter(
            Attendance.employee_id == emp.id,
            extract('month', Attendance.date) == month,
            extract('year', Attendance.date) == year
        ).all()
        records_dict = {rec.date: rec for rec in records}
        ws.cell(row=row_num, column=1, value=emp.full_name)
        ws.cell(row=row_num, column=2, value=emp.employee_id)
        ws.cell(row=row_num, column=3, value=emp.department or '-')
        summary = {'Present': 0, 'Absent': 0, 'Late': 0, 'Half Day': 0, 'WFH': 0}
        for day in range(1, days_in_month + 1):
            d = date(year, month, day)
            col = day + 3
            cell = ws.cell(row=row_num, column=col)
            if d.weekday() == 6:
                cell.value = 'S'
            elif d in records_dict:
                rec = records_dict[d]
                cell.value = status_short.get(rec.status, rec.status[0])
                if rec.status in summary:
                    summary[rec.status] += 1
            else:
                cell.value = '-'
            cell.alignment = center
        base_col = days_in_month + 4
        ws.cell(row=row_num, column=base_col, value=summary['Present'])
        ws.cell(row=row_num, column=base_col+1, value=summary['Absent'])
        ws.cell(row=row_num, column=base_col+2, value=summary['Late'])
        ws.cell(row=row_num, column=base_col+3, value=summary['Half Day'])
        ws.cell(row=row_num, column=base_col+4, value=summary['WFH'])
        working = sum(1 for d in range(1, days_in_month+1)
                      if date(year, month, d).weekday() != 6)
        pct = round(
            (summary['Present'] + summary['Late'] + summary['WFH']) /
            working * 100, 1
        ) if working else 0
        ws.cell(row=row_num, column=base_col+5, value=f'{pct}%')
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True,
                     download_name=f'attendance_{calendar.month_name[month]}_{year}.xlsx')