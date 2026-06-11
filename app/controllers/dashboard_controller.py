from flask import Blueprint, render_template, redirect, url_for, flash, request
from flask_login import login_required, current_user
from app.models.employee import Employee
from app.models.user import User
from app.models.role import Role
from app.models.mood import MoodTracker
from app.extensions import db
from app.utils.decorators import admin_required
from sqlalchemy import func
from datetime import date

dashboard_bp = Blueprint('dashboard', __name__)


@dashboard_bp.route('/')
@dashboard_bp.route('/index')
@login_required
def index():
    from datetime import date
    today = date.today()

    total_employees = Employee.query.count()
    active_employees = Employee.query.filter_by(status='Active').count()
    inactive_employees = Employee.query.filter_by(status='Inactive').count()

    dept_data = db.session.query(
        Employee.department,
        func.count(Employee.id).label('count')
    ).group_by(Employee.department).all()

    dept_labels = [row[0] or 'Unknown' for row in dept_data]
    dept_counts = [row[1] for row in dept_data]

    type_data = db.session.query(
        Employee.employment_type,
        func.count(Employee.id).label('count')
    ).group_by(Employee.employment_type).all()

    type_labels = [row[0] or 'Unknown' for row in type_data]
    type_counts = [row[1] for row in type_data]

    recent_employees = Employee.query.order_by(
        Employee.created_at.desc()
    ).limit(5).all()

    # Mood summary
    today_mood_summary = {}
    for key in MoodTracker.MOODS:
        count = MoodTracker.query.filter_by(
            date=today, mood=key
        ).count()
        today_mood_summary[key] = count

    # ============ BIRTHDAY LOGIC ============
    all_employees = Employee.query.filter(
        Employee.status == 'Active',
        Employee.date_of_birth != None
    ).all()

    todays_birthdays = []
    upcoming_birthdays = []

    for emp in all_employees:
        if emp.date_of_birth:
            # Aaj ka birthday
            if (emp.date_of_birth.month == today.month and
                    emp.date_of_birth.day == today.day):
                todays_birthdays.append(emp)
            else:
                # Upcoming 7 days
                from datetime import timedelta
                for i in range(1, 8):
                    future = today + timedelta(days=i)
                    if (emp.date_of_birth.month == future.month and
                            emp.date_of_birth.day == future.day):
                        upcoming_birthdays.append({
                            'employee': emp,
                            'days_left': i,
                            'date': future
                        })
                        break

    # Sort upcoming by days
    upcoming_birthdays.sort(key=lambda x: x['days_left'])

    return render_template('dashboard/index.html',
                           total=total_employees,
                           active=active_employees,
                           inactive=inactive_employees,
                           dept_labels=dept_labels,
                           dept_counts=dept_counts,
                           type_labels=type_labels,
                           type_counts=type_counts,
                           recent_employees=recent_employees,
                           today_mood_summary=today_mood_summary,
                           todays_birthdays=todays_birthdays,
                           upcoming_birthdays=upcoming_birthdays,
                           today=today)


@dashboard_bp.route('/users')
@admin_required
def manage_users():
    from app.models.employee import Employee
    users = User.query.all()
    roles = Role.query.all()
    # Unlinked employees
    unlinked_employees = Employee.query.filter_by(
        user_id=None, status='Active'
    ).all()
    return render_template('dashboard/users.html',
                           users=users,
                           roles=roles,
                           unlinked_employees=unlinked_employees)

@dashboard_bp.route('/users/add', methods=['POST'])
@admin_required
def add_user():
    from app.models.employee import Employee
    username = request.form.get('username')
    email = request.form.get('email')
    password = request.form.get('password')
    role_id = request.form.get('role_id') or None
    employee_id = request.form.get('employee_id') or None

    existing = User.query.filter_by(username=username).first()
    if existing:
        flash(f'Username "{username}" already exists!', 'danger')
        return redirect(url_for('dashboard.manage_users'))

    new_user = User(
        username=username,
        email=email,
        active=True,
        role_id=role_id
    )
    new_user.set_password(password)
    db.session.add(new_user)
    db.session.flush()

    # Auto link employee
    if employee_id:
        emp = Employee.query.get(int(employee_id))
        if emp:
            emp.user_id = new_user.id
            flash(f'User "{username}" linked to {emp.full_name}!', 'success')

    db.session.commit()
    flash(f'User "{username}" created!', 'success')
    return redirect(url_for('dashboard.manage_users'))


@dashboard_bp.route('/users/<int:id>/change-role', methods=['POST'])
@admin_required
def change_role(id):
    user = User.query.get_or_404(id)
    role_id = request.form.get('role_id')
    user.role_id = role_id if role_id else None
    db.session.commit()
    flash(f'{user.username} role updated!', 'success')
    return redirect(url_for('dashboard.manage_users'))


@dashboard_bp.route('/users/<int:id>/toggle', methods=['POST'])
@admin_required
def toggle_user(id):
    user = User.query.get_or_404(id)
    if user.username == 'admin':
        flash('Admin account disable nahi kar sakte!', 'danger')
        return redirect(url_for('dashboard.manage_users'))
    user.active = not user.active
    db.session.commit()
    status = 'Activated' if user.active else 'Deactivated'
    flash(f'{user.username} {status}!', 'success')
    return redirect(url_for('dashboard.manage_users'))


@dashboard_bp.route('/users/<int:id>/delete', methods=['POST'])
@admin_required
def delete_user(id):
    user = User.query.get_or_404(id)
    if user.username == 'admin':
        flash('Admin account delete nahi kar sakte!', 'danger')
        return redirect(url_for('dashboard.manage_users'))
    db.session.delete(user)
    db.session.commit()
    flash(f'{user.username} deleted!', 'success')
    return redirect(url_for('dashboard.manage_users'))


@dashboard_bp.route('/users/<int:id>/reset-password', methods=['POST'])
@admin_required
def reset_password(id):
    user = User.query.get_or_404(id)
    new_password = request.form.get('new_password')
    if not new_password or len(new_password) < 6:
        flash('Password minimum 6 characters!', 'danger')
        return redirect(url_for('dashboard.manage_users'))
    user.set_password(new_password)
    db.session.commit()
    flash(f'{user.username} password reset!', 'success')
    return redirect(url_for('dashboard.manage_users'))

@dashboard_bp.route('/users/<int:id>/link-employee', methods=['POST'])
@admin_required
def link_employee(id):
    from app.models.employee import Employee
    user = User.query.get_or_404(id)
    employee_id = request.form.get('employee_id')

    if not employee_id:
        flash('Please select an employee!', 'danger')
        return redirect(url_for('dashboard.manage_users'))

    emp = Employee.query.get(int(employee_id))
    if emp:
        emp.user_id = user.id
        db.session.commit()
        flash(f' {user.username} linked to {emp.full_name}!', 'success')
    return redirect(url_for('dashboard.manage_users'))


@dashboard_bp.route('/export/excel')
@login_required
def export_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    import io
    from flask import send_file

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees"

    dark = PatternFill(start_color='1a1a2e', end_color='1a1a2e', fill_type='solid')
    white_font = Font(color='FFFFFF', bold=True)
    center = Alignment(horizontal='center')

    headers = ['Employee ID', 'First Name', 'Last Name', 'Email',
               'Phone', 'Department', 'Designation', 'Employment Type',
               'Date Joined', 'Salary', 'Status', 'Gender', 'Blood Group']

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = dark
        cell.font = white_font
        cell.alignment = center

    employees = Employee.query.order_by(Employee.employee_id).all()
    for row, emp in enumerate(employees, 2):
        ws.cell(row=row, column=1, value=emp.employee_id)
        ws.cell(row=row, column=2, value=emp.first_name)
        ws.cell(row=row, column=3, value=emp.last_name)
        ws.cell(row=row, column=4, value=emp.email)
        ws.cell(row=row, column=5, value=emp.phone or '')
        ws.cell(row=row, column=6, value=emp.department or '')
        ws.cell(row=row, column=7, value=emp.designation or '')
        ws.cell(row=row, column=8, value=emp.employment_type or '')
        ws.cell(row=row, column=9, value=str(emp.date_joined) if emp.date_joined else '')
        ws.cell(row=row, column=10, value=float(emp.salary) if emp.salary else '')
        ws.cell(row=row, column=11, value=emp.status)
        ws.cell(row=row, column=12, value=emp.gender or '')
        ws.cell(row=row, column=13, value=emp.blood_group or '')

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='employees.xlsx'
    )