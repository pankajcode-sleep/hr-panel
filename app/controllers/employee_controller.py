from flask import Blueprint, render_template, redirect, url_for, flash, request, send_file
from flask_login import login_required, current_user
from app.extensions import db
from app.models.employee import Employee
from app.utils.decorators import admin_required, hr_required
from app.services.document_service import generate_payslip, generate_offer_letter, generate_increment_letter
import os
from werkzeug.utils import secure_filename

employee_bp = Blueprint('employees', __name__)

UPLOAD_FOLDER = os.path.join('app', 'static', 'uploads', 'photos')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif'}


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def save_photo(file, employee_id):
    if file and allowed_file(file.filename):
        ext = file.filename.rsplit('.', 1)[1].lower()
        import time
        filename = f'emp_{employee_id}_{int(time.time())}.{ext}'
        file.save(os.path.join(UPLOAD_FOLDER, filename))
        return filename
    return None


def generate_employee_id():
    last_employee = Employee.query.order_by(Employee.id.desc()).first()
    if not last_employee:
        return 'EMP-001'
    try:
        number = int(last_employee.employee_id.split('-')[1])
        return f'EMP-{(number + 1):03d}'
    except:
        count = Employee.query.count()
        return f'EMP-{(count + 1):03d}'


@employee_bp.route('/')
@login_required
def list_employees():
    search = request.args.get('q', '')
    page = request.args.get('page', 1, type=int)
    query = Employee.query
    if search:
        query = query.filter(
            db.or_(
                Employee.first_name.ilike(f'%{search}%'),
                Employee.last_name.ilike(f'%{search}%'),
                Employee.email.ilike(f'%{search}%'),
                Employee.department.ilike(f'%{search}%'),
            )
        )
    employees = query.order_by(Employee.created_at.desc()).paginate(
        page=page, per_page=10, error_out=False
    )
    return render_template('employee/list.html', employees=employees, search=search)


@employee_bp.route('/add', methods=['GET', 'POST'])
@hr_required
def add_employee():
    next_id = generate_employee_id()
    if request.method == 'POST':
        employee = Employee(
            first_name=request.form['first_name'],
            last_name=request.form['last_name'],
            email=request.form['email'],
            employee_id=request.form['employee_id'],
            department=request.form.get('department'),
            designation=request.form.get('designation'),
            employment_type=request.form.get('employment_type', 'Full-time'),
            phone=request.form.get('phone'),
            date_joined=request.form.get('date_joined') or None,
            date_of_birth=request.form.get('date_of_birth') or None,
            gender=request.form.get('gender'),
            blood_group=request.form.get('blood_group'),
            marital_status=request.form.get('marital_status'),
            nationality=request.form.get('nationality'),
            religion=request.form.get('religion'),
            emergency_contact=request.form.get('emergency_contact'),
            salary=request.form.get('salary') or None,
            address=request.form.get('address'),
            status='Active'
        )
        db.session.add(employee)
        db.session.flush()
        photo = request.files.get('photo')
        if photo and photo.filename:
            filename = save_photo(photo, employee.employee_id)
            if filename:
                employee.photo = filename
        db.session.commit()
        flash('Employee added successfully!', 'success')
        return redirect(url_for('employees.list_employees'))
    return render_template('employee/add.html', next_id=next_id)


@employee_bp.route('/<int:id>')
@login_required
def view_employee(id):
    employee = Employee.query.get_or_404(id)
    return render_template('employee/view.html', employee=employee)


@employee_bp.route('/<int:id>/edit', methods=['GET', 'POST'])
@hr_required
def edit_employee(id):
    employee = Employee.query.get_or_404(id)
    if request.method == 'POST':
        employee.first_name = request.form['first_name']
        employee.last_name = request.form['last_name']
        employee.email = request.form['email']
        employee.employee_id = request.form['employee_id']
        employee.department = request.form.get('department')
        employee.designation = request.form.get('designation')
        employee.employment_type = request.form.get('employment_type', 'Full-time')
        employee.phone = request.form.get('phone')
        employee.date_joined = request.form.get('date_joined') or None
        employee.date_of_birth = request.form.get('date_of_birth') or None
        employee.gender = request.form.get('gender')
        employee.blood_group = request.form.get('blood_group')
        employee.marital_status = request.form.get('marital_status')
        employee.nationality = request.form.get('nationality')
        employee.religion = request.form.get('religion')
        employee.emergency_contact = request.form.get('emergency_contact')
        employee.status = request.form.get('status', 'Active')
        employee.salary = request.form.get('salary') or None
        employee.address = request.form.get('address')

        photo = request.files.get('photo')
        if photo and photo.filename:
            if employee.photo:
                old_path = os.path.join(UPLOAD_FOLDER, employee.photo)
                if os.path.exists(old_path):
                    os.remove(old_path)
            filename = save_photo(photo, employee.employee_id)
            if filename:
                employee.photo = filename

        db.session.commit()
        flash(f'{employee.full_name} updated!', 'success')
        return redirect(url_for('employees.view_employee', id=employee.id))
    return render_template('employee/edit.html', employee=employee)


@employee_bp.route('/<int:id>/delete', methods=['POST'])
@admin_required
def delete_employee(id):
    employee = Employee.query.get_or_404(id)
    try:
        from app.models.attendance import Attendance
        from app.models.leave import LeaveApplication, LeaveBalance, CompOff
        from app.models.mood import MoodTracker
        Attendance.query.filter_by(employee_id=id).delete()
        LeaveApplication.query.filter_by(employee_id=id).delete()
        LeaveBalance.query.filter_by(employee_id=id).delete()
        CompOff.query.filter_by(employee_id=id).delete()
        MoodTracker.query.filter_by(employee_id=id).delete()
        db.session.delete(employee)
        db.session.commit()
        flash(f'{employee.full_name} deleted!', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Error: {str(e)}', 'danger')
    return redirect(url_for('employees.list_employees'))


@employee_bp.route('/<int:id>/toggle-status', methods=['POST'])
@hr_required
def toggle_status(id):
    employee = Employee.query.get_or_404(id)
    employee.status = 'Inactive' if employee.status == 'Active' else 'Active'
    db.session.commit()
    flash(f'{employee.full_name} status changed to {employee.status}!', 'success')
    return redirect(url_for('employees.list_employees'))


@employee_bp.route('/<int:id>/id-card')
@login_required
def id_card(id):
    employee = Employee.query.get_or_404(id)
    return render_template('employee/id_card.html', employee=employee)


@employee_bp.route('/<int:id>/payslip', methods=['GET', 'POST'])
@hr_required
def payslip(id):
    employee = Employee.query.get_or_404(id)
    if request.method == 'POST':
        month = int(request.form.get('month'))
        year = int(request.form.get('year'))
        buffer = generate_payslip(employee, month, year)
        return send_file(buffer, mimetype='application/pdf',
                         as_attachment=True,
                         download_name=f'payslip_{employee.employee_id}_{month}_{year}.pdf')
    return render_template('employee/payslip.html', employee=employee)


@employee_bp.route('/<int:id>/offer-letter')
@hr_required
def offer_letter(id):
    employee = Employee.query.get_or_404(id)
    buffer = generate_offer_letter(employee)
    return send_file(buffer, mimetype='application/pdf',
                     as_attachment=True,
                     download_name=f'offer_letter_{employee.employee_id}.pdf')


@employee_bp.route('/<int:id>/increment-letter', methods=['GET', 'POST'])
@hr_required
def increment_letter(id):
    employee = Employee.query.get_or_404(id)
    if request.method == 'POST':
        old_salary = request.form.get('old_salary')
        new_salary = request.form.get('new_salary')
        effective_date = request.form.get('effective_date')
        buffer = generate_increment_letter(employee, old_salary, new_salary, effective_date)
        return send_file(buffer, mimetype='application/pdf',
                         as_attachment=True,
                         download_name=f'increment_{employee.employee_id}.pdf')
    return render_template('employee/increment_letter.html', employee=employee)


@employee_bp.route('/export/excel')
@login_required
def export_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    import io

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees"
    dark = PatternFill(start_color='1a1a2e', end_color='1a1a2e', fill_type='solid')
    white_font = Font(color='FFFFFF', bold=True)
    center = Alignment(horizontal='center')
    headers = ['Employee ID', 'First Name', 'Last Name', 'Email',
               'Phone', 'Department', 'Designation', 'Employment Type',
               'Date Joined', 'Salary', 'Status', 'Gender', 'Blood Group']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = dark
        cell.font = white_font
        cell.alignment = center
    employees = Employee.query.order_by(Employee.employee_id).all()
    for row, emp in enumerate(employees, 2):
        ws.cell(row=row, column=1, value=emp.employee_id)
        ws.cell(row=row, column=2, value=emp.first_name)
        ws.cell(row=row, column=3, value=emp.last_name)
        ws.cell(row=row, column=4, value=emp.email)
        ws.cell(row=row, column=5, value=emp.phone or '')
        ws.cell(row=row, column=6, value=emp.department or '')
        ws.cell(row=row, column=7, value=emp.designation or '')
        ws.cell(row=row, column=8, value=emp.employment_type or '')
        ws.cell(row=row, column=9, value=str(emp.date_joined) if emp.date_joined else '')
        ws.cell(row=row, column=10, value=float(emp.salary) if emp.salary else '')
        ws.cell(row=row, column=11, value=emp.status)
        ws.cell(row=row, column=12, value=emp.gender or '')
        ws.cell(row=row, column=13, value=emp.blood_group or '')
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True,
                     download_name='employees.xlsx')